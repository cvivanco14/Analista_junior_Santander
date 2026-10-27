"""Arma las tablas de resultados y la base en Excel con todo lo utilizado.

Salida: output/Base_Proyeccion_IPC.xlsx y output/tablas/*.csv (insumo del reporte HTML)
"""
import os

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import norm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "output")
TAB = os.path.join(OUT, "tablas")
INS = os.path.join(ROOT, "data", "insumos")
os.makedirs(TAB, exist_ok=True)

base = pd.read_csv(os.path.join(ROOT, "data", "base_modelo.csv"), index_col=0, parse_dates=True)
pw = pd.read_csv(os.path.join(ROOT, "data", "pesos_bloques.csv"))
W = pw.set_index("bloque").peso_2023
HOJAS = list(W.index)
ULT = base["idx_" + HOJAS[0]].dropna().index.max()
CENTRAL = "Combinado"

AGREGADOS = {
    "IPC": HOJAS,
    "IPC sin volátiles": ["Alim_SV", "Bs_sv_ex_alim", "ss_sv_ai", "ss_sv_ex_ai"],
    "Bienes SV": ["Alim_SV", "Bs_sv_ex_alim"],
    "Servicios SV": ["ss_sv_ai", "ss_sv_ex_ai"],
    "IPC volátiles": ["elec", "comb", "alim_rest_vol", "fyv", "bs_vol", "ss_v_ai", "ss_v_ex_ai"],
    "Energía volátil": ["elec", "comb"],
    "Alimentos volátiles": ["alim_rest_vol", "fyv"],
    "Resto volátiles": ["bs_vol", "ss_v_ai", "ss_v_ex_ai"],
}
NOMBRE = dict(zip(pw.bloque, pw.nombre))

# ------------------------------------------------------------------ niveles proyectados
pm = pd.read_csv(os.path.join(OUT, "proy_bloques_modelos.csv"), parse_dates=["periodo"])
pm["indice"] = np.exp(pm.log_indice / 100)
hist = base[["idx_" + b for b in HOJAS]].rename(columns=lambda c: c[4:]).dropna()


def niveles(modelo):
    q = pm[pm.modelo == modelo].pivot(index="periodo", columns="bloque", values="indice")[HOJAS]
    return pd.concat([hist.loc[: q.index.min() - pd.DateOffset(days=1)], q])


def agregar(lv):
    return pd.DataFrame({a: (lv[c] * W[c]).sum(axis=1) / W[c].sum() for a, c in AGREGADOS.items()})


modelos = pm.modelo.unique().tolist()
NIV = {m: niveles(m) for m in modelos}
AGR = {m: agregar(NIV[m]) for m in modelos}
fut = AGR[CENTRAL].loc[ULT + pd.DateOffset(months=1):].index

lv_c, ag_c = NIV[CENTRAL], AGR[CENTRAL]
mm = ag_c.pct_change() * 100
a12 = ag_c.pct_change(12) * 100
dic_prev = ag_c[ag_c.index.month == 12]
acum = pd.DataFrame({c: [100 * (ag_c.loc[t, c] / dic_prev.loc[f"{t.year - 1}-12-01", c] - 1)
                         if pd.Timestamp(f"{t.year - 1}-12-01") in dic_prev.index else np.nan
                         for t in ag_c.index] for c in ag_c.columns}, index=ag_c.index)

# ------------------------------------------------------------------ bandas (errores fuera de muestra)
ei = pd.read_csv(os.path.join(OUT, "eval_errores_ipc.csv"), parse_dates=["origen", "objetivo"])
ec = ei[(ei.modelo == CENTRAL)].dropna(subset=["err_12m"])
qs = ec.groupby("h").err_12m.quantile([0.05, 0.10, 0.30, 0.70, 0.90, 0.95]).unstack()
qs_rec = ec[ec.origen >= "2023-07-01"].groupby("h").err_12m.quantile([0.10, 0.30, 0.70, 0.90]).unstack()
bandas = pd.DataFrame(index=fut)
for i, t in enumerate(fut, start=1):
    for q in qs.columns:
        bandas.loc[t, f"p{int(q * 100):02d}"] = a12.loc[t, "IPC"] + qs.loc[i, q]
    for q in qs_rec.columns:
        bandas.loc[t, f"p{int(round(q * 100)):02d}_pos2023"] = a12.loc[t, "IPC"] + qs_rec.loc[i, q]
bandas.insert(0, "central", a12.loc[fut, "IPC"])

# ------------------------------------------------------------------ tabla principal
filas = []
for t in fut:
    r = {"periodo": t.strftime("%Y-%m"), "IPC_indice": ag_c.loc[t, "IPC"],
         "IPC_var_mm": mm.loc[t, "IPC"], "IPC_var_12m": a12.loc[t, "IPC"], "IPC_var_acum": acum.loc[t, "IPC"]}
    for a in list(AGREGADOS)[1:]:
        r[f"{a}_mm"] = mm.loc[t, a]
        r[f"{a}_12m"] = a12.loc[t, a]
    for m in ["M1_AR_est", "M2_ARDL_macro", "M3_ECM_tc"]:
        r[f"IPC_12m_{m}"] = AGR[m].pct_change(12).loc[t, "IPC"] * 100
    filas.append(r)
proy = pd.DataFrame(filas)
proy = proy.join(bandas.reset_index(drop=True).drop(columns="central").add_prefix("banda_12m_"))

# ------------------------------------------------------------------ incidencias y pesos efectivos
ipc_lv = (lv_c * W).sum(axis=1) / 100
inc_mm = lv_c.diff().mul(W / 100).div(ipc_lv.shift(), axis=0) * 100
inc_12 = (lv_c - lv_c.shift(12)).mul(W / 100).div(ipc_lv.shift(12), axis=0) * 100
w_ef = lv_c.shift(1).mul(W).div(ipc_lv.shift(1), axis=0)
inc_tab = pd.DataFrame({
    "bloque": HOJAS, "nombre": [NOMBRE[b] for b in HOJAS], "peso_2023": W.values,
    "peso_efectivo_oct26": w_ef.loc[fut[0]].values,
    "var_12m_dic26": (lv_c.loc["2026-12-01"] / lv_c.loc["2025-12-01"] - 1).values * 100,
    "incid_12m_dic26": inc_12.loc["2026-12-01"].values,
    "var_12m_dic27": (lv_c.loc["2027-12-01"] / lv_c.loc["2026-12-01"] - 1).values * 100,
    "incid_12m_dic27": inc_12.loc["2027-12-01"].values,
})
inc_mm_tab = inc_mm.loc[fut].rename(columns=NOMBRE)
inc_mm_tab.index = inc_mm_tab.index.strftime("%Y-%m")

bloques_mm = (lv_c.pct_change() * 100).loc[fut].rename(columns=NOMBRE)
bloques_mm.index = bloques_mm.index.strftime("%Y-%m")
modelos_mm = []
for m in modelos:
    x = (NIV[m].pct_change() * 100).loc[fut]
    x.index = x.index.strftime("%Y-%m")
    x = x.rename(columns=NOMBRE)
    x.insert(0, "IPC", (AGR[m].pct_change() * 100).loc[fut, "IPC"].values)
    x.insert(0, "modelo", m)
    modelos_mm.append(x.reset_index())
modelos_mm = pd.concat(modelos_mm)

# ------------------------------------------------------------------ evaluación
ev = ei[ei.origen >= "2020-01-01"]
rmse = ev.groupby(["modelo", "h"]).agg(
    rmse_mm=("err_mm", lambda e: np.sqrt(np.mean(e ** 2))),
    rmse_12m=("err_12m", lambda e: np.sqrt(np.mean(e ** 2))),
    sesgo_12m=("err_12m", "mean"), n=("err_mm", "size")).reset_index()
ev_rec = ei[ei.origen >= "2023-07-01"]
rmse_rec = ev_rec.groupby(["modelo", "h"]).agg(
    rmse_mm=("err_mm", lambda e: np.sqrt(np.mean(e ** 2))),
    rmse_12m=("err_12m", lambda e: np.sqrt(np.mean(e ** 2)))).reset_index()


def dm(m1, m2, h, col, ev):
    a = ev[(ev.modelo == m1) & (ev.h == h)].set_index("objetivo")[col]
    b = ev[(ev.modelo == m2) & (ev.h == h)].set_index("objetivo")[col]
    j = pd.concat([a, b], axis=1, keys=["a", "b"]).dropna()
    dd = (j.a ** 2 - j.b ** 2).values
    n = len(dd)
    g = [np.mean((dd[k:] - dd.mean()) * (dd[: n - k] - dd.mean())) for k in range(h)]
    v = (g[0] + 2 * sum((1 - k / h) * g[k] for k in range(1, h))) / n
    s = dd.mean() / np.sqrt(v)
    return s, 2 * (1 - norm.cdf(abs(s))), n


dm_rows = []
for periodo, evx in [("2020-2026", ev), ("jul.2023-2026", ev_rec)]:
    for rival in ["D1_AR_est_agregado", "D2_paseo_aleatorio_12m", "M1_AR_est", "M2_ARDL_macro", "M3_ECM_tc", "EEE"]:
        for h, col in [(1, "err_mm"), (3, "err_12m"), (12, "err_12m")]:
            if rival == "EEE" and h != 1:
                continue
            s, p, n = dm(CENTRAL, rival, h, col, evx)
            dm_rows.append({"muestra_origenes": periodo, "modelo": CENTRAL, "rival": rival, "h": h,
                            "variable": "var. m/m" if col == "err_mm" else "var. 12m",
                            "estadistico_DM": s, "p_valor": p, "n": n})
dm_tab = pd.DataFrame(dm_rows)

# ------------------------------------------------------------------ comparación con mercado y proyección previa
wb = openpyxl.load_workbook(os.path.join(INS, "IPC_2023_Automatizado_v2.xlsx"), read_only=True, data_only=True)
rows = list(wb["Historia"].iter_rows(min_row=3, max_col=23, values_only=True))
hx = pd.DataFrame([r for r in rows[1:] if r[0] is not None], columns=rows[0])
hx.index = pd.to_datetime(hx.pop("mes").astype(str).str[:7])
hx = hx[[c for c in hx.columns if c]].apply(pd.to_numeric, errors="coerce")
prev = hx[["IPC_ref", "IPC_SV"]].pct_change()
prev12 = hx[["IPC_ref", "IPC_SV"]].pct_change(12) * 100
d = pd.read_csv(os.path.join(ROOT, "data", "bde", "bde_diario.csv"), index_col=0, parse_dates=True)
ult_d = d.dropna(subset=["bei_swap_1a"]).iloc[-1]
comp = pd.DataFrame([
    {"indicador": "IPC oct-26, var. m/m", "esta_proyeccion": mm.loc["2026-10-01", "IPC"],
     "EEE_oct26": base.loc["2026-10-01", "eee_mm_mes"], "proyeccion_previa_excel": prev.loc["2026-10-01", "IPC_ref"] * 100},
    {"indicador": "IPC nov-26, var. m/m", "esta_proyeccion": mm.loc["2026-11-01", "IPC"],
     "EEE_oct26": base.loc["2026-10-01", "eee_mm_sig"], "proyeccion_previa_excel": prev.loc["2026-11-01", "IPC_ref"] * 100},
    {"indicador": "IPC dic-26, var. 12m", "esta_proyeccion": a12.loc["2026-12-01", "IPC"],
     "EEE_oct26": base.loc["2026-10-01", "eee_dic26"], "proyeccion_previa_excel": prev12.loc["2026-12-01", "IPC_ref"]},
    {"indicador": "IPC SV dic-26, var. 12m", "esta_proyeccion": a12.loc["2026-12-01", "IPC sin volátiles"],
     "EEE_oct26": base.loc["2026-10-01", "eee_sv_dic26"], "proyeccion_previa_excel": prev12.loc["2026-12-01", "IPC_SV"]},
    {"indicador": "IPC dic-27, var. 12m", "esta_proyeccion": a12.loc["2027-12-01", "IPC"],
     "EEE_oct26": base.loc["2026-10-01", "eee_dic27"], "proyeccion_previa_excel": np.nan},
    {"indicador": "IPC SV dic-27, var. 12m", "esta_proyeccion": a12.loc["2027-12-01", "IPC sin volátiles"],
     "EEE_oct26": base.loc["2026-10-01", "eee_sv_dic27"], "proyeccion_previa_excel": np.nan},
    {"indicador": "IPC sep-27, var. 12m (vs. compensación swap 1 año)", "esta_proyeccion": a12.loc["2027-09-01", "IPC"],
     "EEE_oct26": np.nan, "proyeccion_previa_excel": np.nan, "compensacion_swap_1a": ult_d["bei_swap_1a"]},
])

# ------------------------------------------------------------------ difusión (productos con alza)
rows = list(wb["Índice 2023"].iter_rows(min_row=1, max_row=284, max_col=110, values_only=True))
cols = [c.strftime("%Y-%m") if hasattr(c, "strftime") else c for c in rows[0][:110]]
prod = pd.DataFrame([r for r in rows[1:] if r[0]], columns=cols).set_index("Code")
niv_p = prod.loc[:, "2023-12":"2026-09"].apply(pd.to_numeric, errors="coerce")
dif = (niv_p.pct_change(axis=1) > 0).mean() * 100
dif = dif.iloc[1:].rename("difusion_IPC").to_frame()
dif.index = pd.to_datetime(dif.index)
dh = pd.read_csv(os.path.join(INS, "difusion_historica.csv"))
dh = dh[dh.agregado == "IPC"].pivot(index="mes", columns="estadistico", values="valor")
dif["hist_prom"] = dif.index.month.map(dh["prom"])
dif["hist_min"] = dif.index.month.map(dh["min"])
dif["hist_max"] = dif.index.month.map(dh["max"])

# ------------------------------------------------------------------ validaciones de insumos
hb = pd.read_csv(os.path.join(INS, "historico_bcch_indices.csv"), parse_dates=["periodo"]).set_index("periodo")
mapa = {"Alim_SV": "Alimentos no volátiles", "Bs_sv_ex_alim": "Bienes sv. ex. alim.",
        "ss_sv_ai": "Servicios_A&I_sv", "ss_sv_ex_ai": "Servicios_ Resto_sv", "elec": "Electricidad_v",
        "comb": "Combustible_v", "alim_rest_vol": "Alim. Resto_v", "fyv": "FyV (proxy)_v",
        "bs_vol": "Bienes volát.", "ss_v_ai": "Servicios_A&I_v", "ss_v_ex_ai": "Servicios_Resto_v"}
val = []
for b, c in mapa.items():
    x = pd.concat([hist[b].pct_change(12), hb[c].pct_change(12)], axis=1, sort=True).loc["2010":"2023"].dropna() * 100
    val.append({"bloque": NOMBRE[b], "serie_historico_bcch": c, "corr_var12m_2010_2023": x.corr().iloc[0, 1],
                "dif_media_abs_pp": (x.iloc[:, 0] - x.iloc[:, 1]).abs().mean()})
val = pd.DataFrame(val)
p18 = pd.read_csv(os.path.join(INS, "pesos_canasta_2018.csv")).set_index("agregado").ponderacion_2018
pesos_cmp = pd.DataFrame({
    "agregado": ["IPC sin volátiles", "Bienes SV", "Alimentos no volátiles", "Bienes sv. ex. alim.", "Servicios SV",
                 "IPC volátiles", "Alimentos volátiles", "Energía volátil", "Resto volátiles"],
    "peso_2018": [p18["IPC sin volátiles"], p18["Bienes sin volátiles"], p18["Alimentos no volátiles"],
                  p18["Bienes sv. ex. alim."], p18["Servicios sin volátiles"], p18["IPC volátiles"],
                  p18["Alim. Volát."], p18["Energía Volátiles"], p18["Resto volátiles"]],
    "peso_2023": [W[AGREGADOS["IPC sin volátiles"]].sum(), W[AGREGADOS["Bienes SV"]].sum(), W["Alim_SV"],
                  W["Bs_sv_ex_alim"], W[AGREGADOS["Servicios SV"]].sum(), W[AGREGADOS["IPC volátiles"]].sum(),
                  W[AGREGADOS["Alimentos volátiles"]].sum(), W[AGREGADOS["Energía volátil"]].sum(),
                  W[AGREGADOS["Resto volátiles"]].sum()],
})

# ------------------------------------------------------------------ guardar tablas para el reporte
tablas = {
    "proyeccion": proy, "bandas": bandas.reset_index(names="periodo"), "incidencias": inc_tab,
    "comparacion": comp, "rmse": rmse, "rmse_pos2023": rmse_rec, "dm": dm_tab,
    "difusion": dif.reset_index(names="periodo"), "validacion": val, "pesos_cmp": pesos_cmp,
    "agregados_central": ag_c.reset_index(names="periodo"),
}
for k, v in tablas.items():
    v.to_csv(os.path.join(TAB, f"{k}.csv"), index=False)
for m in modelos:
    AGR[m].pct_change(12).mul(100).reset_index(names="periodo").to_csv(os.path.join(TAB, f"a12_{m}.csv"), index=False)

# ------------------------------------------------------------------ Excel
cat = pd.read_csv(os.path.join(ROOT, "data", "bde", "catalogo_series.csv"))
bde_m = pd.read_csv(os.path.join(ROOT, "data", "bde", "bde_mensual.csv"))
bde_d = pd.read_csv(os.path.join(ROOT, "data", "bde", "bde_diario.csv"))
sup = pd.read_csv(os.path.join(OUT, "supuestos_exogenos.csv"), index_col=0)
sup.index.name = "periodo"
sup.columns = ["Dólar observado ($/US$)", "Petróleo WTI en pesos (índice)", "FAO en pesos (índice)",
               "Remuneraciones nominales (índice)", "Imacec desest. (índice)", "EEE inflación 11 meses (%)",
               "EEE inflación 23 meses (%)"]
hist_out = hist.rename(columns=NOMBRE)
hist_out.index.name = "periodo"
pesos_tab = pw.rename(columns={"bloque": "codigo_bloque"}).assign(
    peso_efectivo_oct26=w_ef.loc[fut[0]].values)
pcomb = pd.read_csv(os.path.join(OUT, "pesos_combinacion.csv"))
pcomb.loc[pcomb.bloque == "elec", "modelo"] = "Regla de indexación semestral"
pcomb["bloque"] = pcomb.bloque.map(NOMBRE)
coef = pd.read_csv(os.path.join(OUT, "coeficientes.csv"))
coef = coef[coef.bloque != "elec"]
coef["bloque"] = coef.bloque.map(NOMBRE)
tc_p = pd.read_csv(os.path.join(OUT, "tendencias_comunes_pares.csv"))
tc_p[["bloque_a", "bloque_b"]] = tc_p[["bloque_a", "bloque_b"]].apply(lambda s: s.map(NOMBRE))
tc_g = pd.read_csv(os.path.join(OUT, "tendencias_comunes_grupos.csv"))
tc_g["comparte_tendencia_con"] = tc_g.comparte_tendencia_con.fillna("").apply(
    lambda s: ", ".join(NOMBRE[x.strip()] for x in s.split(",") if x.strip()))
tc_g["bloque"] = tc_g.bloque.map(NOMBRE)

leeme = pd.DataFrame({"Contenido": [
    "Base de la proyección del IPC (oct-2026 a dic-2027) por agregación de bloques",
    f"Último IPC efectivo: {ULT:%B %Y}. Datos BDE descargados el {pd.Timestamp.today():%d-%m-%Y}.",
    "",
    "Hojas:",
    "Proyeccion_IPC: IPC y agregados (var. m/m, 12m, acumulada), escenarios por modelo y bandas de error.",
    "Proyeccion_bloques_mm: variación mensual proyectada de los 11 bloques (proyección central).",
    "Incidencias_mm / Incidencias_12m: aportes de cada bloque y pesos efectivos.",
    "Proyeccion_por_modelo: variación mensual de cada bloque e IPC según cada modelo.",
    "Comparacion_mercado: contra EEE (oct-26), compensación inflacionaria y la proyección previa del Excel.",
    "Evaluacion_RMSE / Evaluacion_RMSE_pos2023 / Test_DM: evaluación fuera de muestra recursiva (orígenes 2020-01 a 2026-08).",
    "Pesos_combinacion: pesos de cada modelo por bloque (inverso del ECM fuera de muestra del error acumulado).",
    "Tend_comunes_pares / Tend_comunes_grupos: tests ADF pairwise de precios relativos (Carlomagno y Espasa).",
    "Coeficientes: estimaciones de cada ecuación (muestra 2010-01 a 2026-09).",
    "Supuestos_exogenos: trayectorias de dólar, petróleo, FAO, salarios, Imacec y expectativas.",
    "Pesos_bloques / Pesos_2018_vs_2023: ponderadores base 2023 (hoja Matriz) y comparación con la canasta 2018.",
    "Indices_bloques: historia de los 11 bloques (1998-12 en adelante, base 2023=100).",
    "Difusion: % de productos con alza mensual 2024-2026 vs. rango histórico (difusion_historica.csv).",
    "Validacion_insumos: consistencia de los bloques con historico_bcch_indices.csv.",
    "BDE_mensual / BDE_diario / Catalogo_BDE: variables macro descargadas de la BDE del Banco Central.",
    "",
    "Supuestos: dólar, petróleo y FAO quedan en su último valor (octubre para dólar y petróleo, promedio de datos diarios);",
    "remuneraciones e Imacec crecen a su ritmo de los últimos 12 meses; la expectativa a 11 meses de la EEE converge en 12 meses",
    "a la de 23 meses; electricidad se ajusta en enero y julio por la inflación del semestre previo.",
    "Scripts: scripts/01_descarga_bde.py → 02_base.py → 03_modelo.py → 04_resultados.py → 05_reporte_html.py",
]})

hojas_xl = {
    "LEEME": (leeme, False),
    "Proyeccion_IPC": (proy, False),
    "Proyeccion_bloques_mm": (bloques_mm, True),
    "Incidencias_mm": (inc_mm_tab, True),
    "Incidencias_12m": (inc_tab, False),
    "Proyeccion_por_modelo": (modelos_mm.rename(columns={"index": "periodo", "periodo": "periodo"}), False),
    "Comparacion_mercado": (comp, False),
    "Evaluacion_RMSE": (rmse, False),
    "Evaluacion_RMSE_pos2023": (rmse_rec, False),
    "Test_DM": (dm_tab, False),
    "Pesos_combinacion": (pcomb, False),
    "Tend_comunes_pares": (tc_p, False),
    "Tend_comunes_grupos": (tc_g, False),
    "Coeficientes": (coef, False),
    "Supuestos_exogenos": (sup, True),
    "Pesos_bloques": (pesos_tab, False),
    "Pesos_2018_vs_2023": (pesos_cmp, False),
    "Indices_bloques": (hist_out, True),
    "Difusion": (dif, True),
    "Validacion_insumos": (val, False),
    "BDE_mensual": (bde_m, False),
    "BDE_diario": (bde_d, False),
    "Catalogo_BDE": (cat, False),
}
xl = os.path.join(OUT, "Base_Proyeccion_IPC.xlsx")
with pd.ExcelWriter(xl, engine="openpyxl", datetime_format="yyyy-mm") as w:
    for nombre, (df, idx) in hojas_xl.items():
        df.to_excel(w, sheet_name=nombre, index=idx)
wbx = openpyxl.load_workbook(xl)
head = PatternFill("solid", fgColor="DDE7F3")
for ws in wbx.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2" if ws.title != "LEEME" else None
    for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 60)), start=1):
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo * 0.9), 60 if ws.title == "LEEME" else 28)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.000"
wbx["LEEME"].column_dimensions["A"].width = 140
wbx.save(xl)
print("Excel:", xl)
print(proy[["periodo", "IPC_var_mm", "IPC_var_12m", "IPC sin volátiles_12m", "IPC volátiles_12m"]].round(2).to_string(index=False))
print(comp.round(2).to_string(index=False))
print(dm_tab.round(3).to_string(index=False))
